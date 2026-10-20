"""
Exports PDF / Excel de la guérite
  - Véhicules présents dans le local  -> export_presents_pdf / export_presents_excel
  - Historique des entrées            -> export_entrees_pdf  / export_entrees_excel

Même charte que les fiches atelier (pdf_fiches.py) :
COULEUR_PRIMAIRE #1a3a5c, ACCENT #e67e22, en-tête « CENTRE AUTO La PRUDENCE+ »,
bandeau de titre foncé souligné d'orange, encadré compteur à droite.

Les 4 vues gardent leurs noms d'origine : urls.py ne change pas.
"""
from collections import Counter, namedtuple
from datetime import datetime, timedelta
from html import escape
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from reportlab.graphics.shapes import Drawing, Ellipse, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.accounts.choices import MotifEntree, StatutVehicule
from apps.guerite.views import get_entrees_presentes_filtrees, get_filtered_entrees


# ════════════════════════════════════════════════════════════════
#  Charte (identique à pdf_fiches.py)
# ════════════════════════════════════════════════════════════════
# Si tu préfères une seule source : from apps.atelier.pdf_fiches import COULEUR_PRIMAIRE, ...
COULEUR_PRIMAIRE = colors.HexColor('#1a3a5c')
COULEUR_ACCENT = colors.HexColor('#e67e22')
COULEUR_GRIS = colors.HexColor('#f5f5f5')
COULEUR_TEXTE = colors.HexColor('#2c3e50')
COULEUR_BORD = colors.HexColor('#9aa7b4')
COULEUR_DISCRET = colors.HexColor('#6b7280')

SERVICES = ("Tôlerie – Peinture – Freinage – Chaudronnerie – Mécanique Générale<br/>"
            "Électricité – Climatisation – Pneumatique – Tapisserie – Sérigraphie")
ADRESSE = "BP : 9060 Douala – Tél : (+237) 650 99 75 09"

PAGE_W, PAGE_H = landscape(A4)
MARGE = 10 * mm
LARGEUR = PAGE_W - 2 * MARGE  # 277 mm

# Couleur d'un motif / statut, devinée depuis son code (adapte si besoin).
# Ordre important : le premier mot-clé trouvé gagne.
TEINTES = [
    (('REPAR',), '#e67e22'),                       # Réparation : orange
    (('VISIT', 'DIAGN', 'CONTROL'), '#2563eb'),    # Visite : bleu
    (('SORTI', 'LIVR', 'TERMIN'), '#16a34a'),      # Sorti : vert
    (('ATTENTE', 'BLOQU', 'ANNUL'), '#dc2626'),    # En attente / annulé : rouge
    (('PRESENT', 'ATELIER', 'COURS', 'ENTR'), '#1a3a5c'),
]
TEINTE_DEFAUT = '#6b7280'


def _teinte(code):
    s = str(code or '').upper()
    for mots, couleur in TEINTES:
        if any(m in s for m in mots):
            return couleur
    return TEINTE_DEFAUT


# ════════════════════════════════════════════════════════════════
#  Utilitaires de valeurs
# ════════════════════════════════════════════════════════════════
def _get(obj, chemin, defaut=None):
    """_get(e, 'vehicule.client.telephone') sans planter sur un None."""
    for attr in chemin.split('.'):
        if obj is None:
            return defaut
        obj = getattr(obj, attr, None)
    return defaut if obj in (None, '') else obj


def _display(obj, champ):
    fn = getattr(obj, f'get_{champ}_display', None)
    if callable(fn):
        try:
            return str(fn())
        except Exception:
            pass
    return str(_get(obj, champ, ''))


def _local(dt):
    """Datetime local naïf (Excel refuse les dates avec fuseau)."""
    if not dt:
        return None
    if isinstance(dt, datetime) and timezone.is_aware(dt):
        dt = timezone.localtime(dt)
    return dt.replace(tzinfo=None) if isinstance(dt, datetime) else dt


def _date_txt(dt, fmt='%d/%m/%Y %H:%M'):
    dt = _local(dt)
    return dt.strftime(fmt) if dt else ''


def _duree(v):
    """timedelta -> « 2 j 05 h », « 3 h 20 min » ; sinon texte brut."""
    if v in (None, ''):
        return ''
    if isinstance(v, timedelta):
        minutes = int(v.total_seconds() // 60)
        j, reste = divmod(minutes, 1440)
        h, m = divmod(reste, 60)
        if j:
            return f"{j} j {h:02d} h"
        return f"{h} h {m:02d} min" if h else f"{m} min"
    return str(v)


def _nom(v):
    return str(v) if v else ''


def _queryset(entrees):
    """Accepte un QuerySet ou une Page (paginator) : on exporte TOUT le résultat filtré."""
    if hasattr(entrees, 'paginator'):
        entrees = entrees.paginator.object_list
    if hasattr(entrees, 'select_related'):
        entrees = entrees.select_related('vehicule', 'vehicule__client', 'conducteur', 'agent_entree')
    return list(entrees)


def _horodatage():
    return timezone.localtime() if timezone.is_aware(timezone.now()) else timezone.now()


# ════════════════════════════════════════════════════════════════
#  Colonnes : une seule définition pour le PDF et l'Excel
# ════════════════════════════════════════════════════════════════
# genre : 'texte' | 'immat' | 'date' | 'badge'
# valeur(e) -> texte (ou datetime pour 'date', ou (libellé, code) pour 'badge')
Colonne = namedtuple('Colonne', 'titre pdf_mm xlsx genre valeur')

COLONNES_PRESENTS = [
    Colonne('Immatriculation', 30, 17, 'immat', lambda e: _get(e, 'vehicule.immatriculation', '')),
    Colonne('Véhicule', 42, 24, 'texte',
            lambda e: f"{_get(e, 'vehicule.marque', '')} {_get(e, 'vehicule.modele', '')}".strip()),
    Colonne('Propriétaire', 44, 26, 'texte', lambda e: _nom(_get(e, 'vehicule.client'))),
    Colonne('Téléphone', 28, 16, 'texte', lambda e: _get(e, 'vehicule.client.telephone', '')),
    Colonne('Conducteur', 38, 22, 'texte', lambda e: _nom(_get(e, 'conducteur'))),
    Colonne('Motif', 30, 16, 'badge', lambda e: (_display(e, 'motif'), _get(e, 'motif'))),
    Colonne('Entrée le', 32, 18, 'date', lambda e: _get(e, 'date_entree')),
    Colonne('Durée', 33, 14, 'texte', lambda e: _duree(_get(e, 'duree_sejour'))),
]

COLONNES_HISTORIQUE = [
    Colonne('N°', 25, 14, 'texte', lambda e: str(_get(e, 'numero', ''))),
    Colonne('Immatriculation', 27, 17, 'immat', lambda e: _get(e, 'vehicule.immatriculation', '')),
    Colonne('Client', 37, 26, 'texte', lambda e: _nom(_get(e, 'vehicule.client'))),
    Colonne('Téléphone', 23, 16, 'texte', lambda e: _get(e, 'vehicule.client.telephone', '')),
    Colonne('Conducteur', 30, 22, 'texte', lambda e: _nom(_get(e, 'conducteur'))),
    Colonne('Motif', 25, 16, 'badge', lambda e: (_display(e, 'motif'), _get(e, 'motif'))),
    Colonne('Statut', 25, 16, 'badge', lambda e: (_display(e, 'statut'), _get(e, 'statut'))),
    Colonne('Entrée', 27, 17, 'date', lambda e: _get(e, 'date_entree')),
    Colonne('Sortie', 27, 17, 'date', lambda e: _get(e, 'date_sortie')),
    Colonne('Agent', 30, 22, 'texte', lambda e: _get(e, 'agent_entree.full_name', '')),
]
# L'Excel de l'historique garde en plus les observations (trop longues pour le PDF).
COLONNES_HISTORIQUE_XLSX = COLONNES_HISTORIQUE + [
    Colonne('Observations', 0, 45, 'texte', lambda e: _get(e, 'observations', '')),
]


def _resume(entrees, champ):
    """[('Réparation', 8, '#e67e22'), ...] trié par effectif décroissant."""
    compte, libelles = Counter(), {}
    for e in entrees:
        code = _get(e, champ)
        compte[code] += 1
        libelles[code] = _display(e, champ) or 'Non renseigné'
    return [(libelles[c], n, _teinte(c)) for c, n in compte.most_common()]


def _filtres_appliques(request):
    g = request.GET
    filtres = []
    if g.get('q'):
        filtres.append(('Recherche', g['q']))
    if g.get('statut'):
        filtres.append(('Statut', dict(StatutVehicule.choices).get(g['statut'], g['statut'])))
    if g.get('motif'):
        filtres.append(('Motif', dict(MotifEntree.choices).get(g['motif'], g['motif'])))
    if g.get('date_debut'):
        filtres.append(('Du', g['date_debut']))
    if g.get('date_fin'):
        filtres.append(('Au', g['date_fin']))
    return filtres


# ════════════════════════════════════════════════════════════════
#  PDF : briques
# ════════════════════════════════════════════════════════════════
def _style(size=7.5, bold=False, align=TA_LEFT, color=COULEUR_TEXTE):
    return ParagraphStyle('', fontName='Helvetica-Bold' if bold else 'Helvetica', fontSize=size,
                          leading=size * 1.2, alignment=align, textColor=color)


def P(texte, size=7.5, bold=False, align=TA_LEFT, color=COULEUR_TEXTE):
    return Paragraph(texte or '', _style(size, bold, align, color))


def _sans_marges(t, valign='MIDDLE'):
    t.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), valign),
        ('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 0),
        ('TOPPADDING', (0, 0), (-1, -1), 0), ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
    ]))
    return t


def _logo(largeur=62 * mm, hauteur=21 * mm):
    """Même logo dessiné que les fiches atelier."""
    d = Drawing(largeur, hauteur)
    d.add(String(largeur / 2, hauteur - 6.5 * mm, "CENTRE AUTO", fontName='Helvetica-Bold',
                 fontSize=11, fillColor=COULEUR_PRIMAIRE, textAnchor='middle'))
    d.add(Ellipse(largeur / 2, hauteur / 2 - 3.2 * mm, largeur / 2 - 1, 6.4 * mm,
                  fillColor=COULEUR_PRIMAIRE, strokeColor=COULEUR_ACCENT, strokeWidth=1.2))
    d.add(String(largeur / 2, hauteur / 2 - 2.6 * mm, "La PRUDENCE+", fontName='Helvetica-Bold',
                 fontSize=14, fillColor=colors.white, textAnchor='middle'))
    d.add(String(largeur / 2, hauteur / 2 - 6.6 * mm, "Véhicules Légers / Poids Lourds",
                 fontName='Helvetica-Oblique', fontSize=6.5, fillColor=colors.white,
                 textAnchor='middle'))
    return d


def _entete(titre, total, unite):
    """Logo + services, puis bandeau titre foncé (souligné orange) + encadré compteur."""
    haut = _sans_marges(Table([[
        _logo(),
        P(f"<b>{SERVICES}</b><br/><font color='#e67e22'><b>{ADRESSE}</b></font>",
          8.5, False, TA_CENTER, COULEUR_PRIMAIRE),
    ]], colWidths=[66 * mm, LARGEUR - 66 * mm]))

    titre_t = Table([[P(f"<b>{escape(titre)}</b>", 13, True, TA_LEFT, colors.white)]],
                    colWidths=[150 * mm])
    titre_t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), COULEUR_PRIMAIRE),
        ('LINEBELOW', (0, 0), (-1, -1), 2, COULEUR_ACCENT),
        ('LEFTPADDING', (0, 0), (-1, -1), 7),
        ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('ROUNDEDCORNERS', [4, 4, 0, 0]),
    ]))
    compteur = Table([[[
        P(f"<font color='#c0392b'><b>{total}</b></font>&nbsp;&nbsp;"
          f"<font size=9 color='#2c3e50'><b>{unite}</b></font>", 15, False, TA_CENTER),
        P(f"Édité le {_horodatage().strftime('%d/%m/%Y à %H:%M')}", 6.5, False, TA_CENTER, COULEUR_DISCRET),
    ]]], colWidths=[62 * mm])
    compteur.setStyle(TableStyle([
        ('BOX', (0, 0), (-1, -1), 1, COULEUR_PRIMAIRE),
        ('ROUNDEDCORNERS', [5, 5, 5, 5]),
        ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    bas = _sans_marges(Table([[titre_t, '', compteur]],
                             colWidths=[150 * mm, LARGEUR - 212 * mm, 62 * mm]))
    return [haut, Spacer(1, 3 * mm), bas, Spacer(1, 3 * mm)]


def _bandeau_infos(resume, filtres):
    """Ligne « Répartition : • Réparation 8 • Visite 4 » + filtres appliqués."""
    morceaux = []
    if resume:
        puces = '&nbsp;&nbsp;&nbsp;'.join(
            f"<font size=9 color='{c}'>•</font>&nbsp;{escape(lib)}&nbsp;<b>{n}</b>" for lib, n, c in resume)
        morceaux.append(f"<font color='#1a3a5c'><b>Répartition :</b></font>&nbsp;&nbsp;{puces}")
    if filtres:
        txt = '&nbsp;&nbsp;·&nbsp;&nbsp;'.join(f"{escape(k)} : <b>{escape(str(v))}</b>" for k, v in filtres)
        morceaux.append(f"<font color='#1a3a5c'><b>Filtres :</b></font>&nbsp;&nbsp;{txt}")
    if not morceaux:
        return []
    t = Table([[P('<br/>'.join(morceaux), 8, False, TA_LEFT)]], colWidths=[LARGEUR])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), COULEUR_GRIS),
        ('LINEBEFORE', (0, 0), (0, -1), 3, COULEUR_ACCENT),
        ('LEFTPADDING', (0, 0), (-1, -1), 7),
        ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    return [t, Spacer(1, 3 * mm)]


def _cellule_pdf(col, e):
    v = col.valeur(e)
    if col.genre == 'immat':
        return P(f"<b>{escape(str(v))}</b>" if v else '—', 8, True, TA_CENTER, COULEUR_PRIMAIRE)
    if col.genre == 'badge':
        libelle, code = v
        if not libelle:
            return P('—', align=TA_CENTER, color=COULEUR_DISCRET)
        return P(f"<font size=9 color='{_teinte(code)}'>•</font>&nbsp;{escape(libelle)}", 7.5, True,
                 TA_LEFT, colors.HexColor(_teinte(code)))
    if col.genre == 'date':
        return P(_date_txt(v) or '—', 7.5, False, TA_CENTER,
                 COULEUR_TEXTE if v else COULEUR_DISCRET)
    return P(escape(str(v)) if v else '—', 7.5, False, TA_LEFT,
             COULEUR_TEXTE if v else COULEUR_DISCRET)


def _tableau_pdf(colonnes, entrees):
    largeurs = [c.pdf_mm for c in colonnes]
    k = LARGEUR / (sum(largeurs) * mm)  # ajuste pile à la largeur utile
    data = [[P(f"<b>{escape(c.titre)}</b>", 8, True, TA_CENTER, colors.white) for c in colonnes]]
    data += [[_cellule_pdf(c, e) for c in colonnes] for e in entrees]
    t = Table(data, colWidths=[w * mm * k for w in largeurs], repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), COULEUR_PRIMAIRE),
        ('LINEBELOW', (0, 0), (-1, 0), 1.5, COULEUR_ACCENT),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, COULEUR_GRIS]),
        ('LINEBELOW', (0, 1), (-1, -1), 0.3, colors.HexColor('#d5dbe1')),
        ('BOX', (0, 0), (-1, -1), 0.8, COULEUR_PRIMAIRE),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, 0), 5), ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('TOPPADDING', (0, 1), (-1, -1), 3.5), ('BOTTOMPADDING', (0, 1), (-1, -1), 3.5),
        ('LEFTPADDING', (0, 0), (-1, -1), 4), ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]))
    return t


def _vide(message):
    t = Table([[P(message, 9.5, True, TA_CENTER, COULEUR_DISCRET)]], colWidths=[LARGEUR])
    t.setStyle(TableStyle([
        ('BOX', (0, 0), (-1, -1), 0.8, COULEUR_BORD),
        ('BACKGROUND', (0, 0), (-1, -1), COULEUR_GRIS),
        ('TOPPADDING', (0, 0), (-1, -1), 14), ('BOTTOMPADDING', (0, 0), (-1, -1), 14),
    ]))
    return t


class _CanvasNumerote(rl_canvas.Canvas):
    """Canvas qui connaît le nombre total de pages (« Page 2 / 5 »)."""
    titre = ''

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._pages = []

    def showPage(self):
        self._pages.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._pages)
        for etat in self._pages:
            self.__dict__.update(etat)
            self._pied(total)
            super().showPage()
        super().save()

    def _pied(self, total):
        self.saveState()
        self.setStrokeColor(COULEUR_ACCENT)
        self.setLineWidth(0.8)
        self.line(MARGE, 9 * mm, PAGE_W - MARGE, 9 * mm)
        self.setFont('Helvetica', 6.5)
        self.setFillColor(COULEUR_DISCRET)
        self.drawString(MARGE, 5.5 * mm,
                        f"Centre Auto La Prudence+ – {self.titre} – imprimé le "
                        f"{_horodatage().strftime('%d/%m/%Y %H:%M')}")
        self.drawRightString(PAGE_W - MARGE, 5.5 * mm, f"Page {self._pageNumber} / {total}")
        self.restoreState()


def _generer_pdf(titre, unite, colonnes, entrees, resume, filtres, message_vide):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4), title=titre, author="Centre Auto La Prudence+",
        topMargin=8 * mm, bottomMargin=13 * mm, leftMargin=MARGE, rightMargin=MARGE,
    )
    elements = _entete(titre.upper(), len(entrees), unite)
    elements += _bandeau_infos(resume, filtres)
    elements.append(_tableau_pdf(colonnes, entrees) if entrees else _vide(message_vide))

    canvas_cls = type('Canvas', (_CanvasNumerote,), {'titre': titre})
    doc.build(elements, canvasmaker=canvas_cls)
    return buffer.getvalue()


# ════════════════════════════════════════════════════════════════
#  Excel
# ════════════════════════════════════════════════════════════════
XL_PRIMAIRE, XL_ACCENT, XL_GRIS, XL_BORD = '1A3A5C', 'E67E22', 'F5F5F5', 'D5DBE1'
_fond = lambda c: PatternFill('solid', fgColor=c)  # noqa: E731
_trait = Side(style='thin', color=XL_BORD)
_BORDURE = Border(left=_trait, right=_trait, top=_trait, bottom=_trait)


def _generer_excel(titre, feuille, colonnes, entrees, resume, filtres):
    wb = Workbook()
    ws = wb.active
    ws.title = feuille[:31]
    ws.sheet_view.showGridLines = False
    n = len(colonnes)
    der = get_column_letter(n)

    def bande(ligne, valeur, font, fill=None, hauteur=None, align='left', bas=None):
        ws.merge_cells(f"A{ligne}:{der}{ligne}")
        c = ws[f"A{ligne}"]
        c.value, c.font = valeur, font
        c.alignment = Alignment(horizontal=align, vertical='center', indent=1, wrap_text=True)
        if fill or bas:
            for col in range(1, n + 1):
                cell = ws.cell(row=ligne, column=col)
                if fill:
                    cell.fill = fill
                if bas:
                    cell.border = Border(bottom=bas)
        if hauteur:
            ws.row_dimensions[ligne].height = hauteur

    # 1-2 : en-tête société (bandeau foncé)
    bande(1, "CENTRE AUTO  La PRUDENCE+", Font(size=16, bold=True, color='FFFFFF'),
          _fond(XL_PRIMAIRE), 28)
    bande(2, f"Véhicules Légers / Poids Lourds   •   {ADRESSE}",
          Font(size=9, bold=True, color=XL_ACCENT), _fond(XL_PRIMAIRE), 18,
          bas=Side(style='thick', color=XL_ACCENT))
    # 4 : titre du document
    bande(4, titre.upper(), Font(size=14, bold=True, color=XL_PRIMAIRE), hauteur=24,
          bas=Side(style='medium', color=XL_ACCENT))
    # 5 : méta
    bande(5, f"Édité le {_horodatage().strftime('%d/%m/%Y à %H:%M')}   •   "
             f"{len(entrees)} enregistrement(s)",
          Font(size=9, italic=True, color='6B7280'), hauteur=16)
    ligne = 6
    if resume:
        bande(ligne, "Répartition :  " + "   •   ".join(f"{lib} : {nb}" for lib, nb, _ in resume),
              Font(size=9, color='2C3E50'), _fond(XL_GRIS), 16)
        ligne += 1
    if filtres:
        bande(ligne, "Filtres :  " + "   •   ".join(f"{k} : {v}" for k, v in filtres),
              Font(size=9, color='2C3E50'), _fond(XL_GRIS), 16)
        ligne += 1

    # En-têtes du tableau
    entete = ligne + 1
    for i, col in enumerate(colonnes, start=1):
        c = ws.cell(row=entete, column=i, value=col.titre)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = _fond(XL_PRIMAIRE)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = Border(left=_trait, right=_trait, bottom=Side(style='medium', color=XL_ACCENT))
        ws.column_dimensions[get_column_letter(i)].width = col.xlsx
    ws.row_dimensions[entete].height = 22

    # Données
    for r, e in enumerate(entrees, start=1):
        row = entete + r
        zebre = _fond(XL_GRIS) if r % 2 == 0 else None
        for i, col in enumerate(colonnes, start=1):
            v = col.valeur(e)
            c = ws.cell(row=row, column=i)
            c.border = _BORDURE
            c.alignment = Alignment(vertical='center', wrap_text=col.genre == 'texte')
            if zebre:
                c.fill = zebre
            if col.genre == 'date':
                c.value = _local(v)
                c.number_format = 'dd/mm/yyyy hh:mm'
                c.alignment = Alignment(horizontal='center', vertical='center')
            elif col.genre == 'badge':
                libelle, code = v
                c.value = libelle
                c.font = Font(bold=True, color=_teinte(code).lstrip('#').upper())
            elif col.genre == 'immat':
                c.value = v
                c.font = Font(bold=True, color=XL_PRIMAIRE)
                c.alignment = Alignment(horizontal='center', vertical='center')
            else:
                c.value = v or None

    fin = entete + max(len(entrees), 1)
    if not entrees:
        ws.merge_cells(f"A{fin}:{der}{fin}")
        c = ws[f"A{fin}"]
        c.value = "Aucun enregistrement."
        c.font = Font(italic=True, color='6B7280')
        c.alignment = Alignment(horizontal='center')

    ws.freeze_panes = f"A{entete + 1}"
    ws.auto_filter.ref = f"A{entete}:{der}{fin}"

    # Impression : paysage, tient en largeur, en-têtes répétés
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = f"{entete}:{entete}"
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.left.text = f"Centre Auto La Prudence+ – {titre}"
    ws.oddFooter.right.text = "Page &P / &N"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ════════════════════════════════════════════════════════════════
#  Vues
# ════════════════════════════════════════════════════════════════
MIME_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _reponse(contenu, content_type, nom):
    resp = HttpResponse(contenu, content_type=content_type)
    resp['Content-Disposition'] = f'attachment; filename="{nom}"'
    return resp


def _suffixe():
    return _horodatage().strftime('%Y%m%d_%H%M')


def export_presents_pdf(request):
    entrees = _queryset(get_entrees_presentes_filtrees(request))
    pdf = _generer_pdf(
        "Véhicules présents dans le local", "véhicule(s)", COLONNES_PRESENTS, entrees,
        _resume(entrees, 'motif'), _filtres_appliques(request),
        "Aucun véhicule dans le local pour le moment.",
    )
    return _reponse(pdf, 'application/pdf', f"vehicules_presents_{_suffixe()}.pdf")


def export_presents_excel(request):
    entrees = _queryset(get_entrees_presentes_filtrees(request))
    xlsx = _generer_excel(
        "Véhicules présents dans le local", "Véhicules présents", COLONNES_PRESENTS, entrees,
        _resume(entrees, 'motif'), _filtres_appliques(request),
    )
    return _reponse(xlsx, MIME_XLSX, f"vehicules_presents_{_suffixe()}.xlsx")


def export_entrees_pdf(request):
    entrees = _queryset(get_filtered_entrees(request))
    pdf = _generer_pdf(
        "Historique des entrées", "mouvement(s)", COLONNES_HISTORIQUE, entrees,
        _resume(entrees, 'statut'), _filtres_appliques(request),
        "Aucun mouvement ne correspond à ces critères.",
    )
    return _reponse(pdf, 'application/pdf', f"historique_entrees_{_suffixe()}.pdf")


def export_entrees_excel(request):
    entrees = _queryset(get_filtered_entrees(request))
    xlsx = _generer_excel(
        "Historique des entrées", "Historique entrées", COLONNES_HISTORIQUE_XLSX, entrees,
        _resume(entrees, 'statut'), _filtres_appliques(request),
    )
    return _reponse(xlsx, MIME_XLSX, f"historique_entrees_{_suffixe()}.xlsx")
